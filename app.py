import os
import re
import sys
import shutil
import socket
import datetime
import sqlite3
import subprocess
import threading
from pathlib import Path
from flask import Flask, render_template, request, jsonify, send_file, send_from_directory
from werkzeug.utils import secure_filename
from pypdf import PdfReader
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "datos_seguros.db"
COMPROBANTES_DIR = BASE_DIR / "comprobantes"
COMPROBANTES_DIR.mkdir(exist_ok=True)
CLOUDFLARED_EXE = BASE_DIR / "cloudflared.exe"

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 32 * 1024 * 1024  # 32 MB max

ENLACE_PUBLICO_ACTUAL = None

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    with get_db() as conn:
        # Tabla de configuración
        conn.execute('''
            CREATE TABLE IF NOT EXISTS config (
                clave TEXT PRIMARY KEY,
                valor TEXT NOT NULL
            )
        ''')

        # Tabla de usuarios y accesos
        conn.execute('''
            CREATE TABLE IF NOT EXISTS usuarios (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                nombre TEXT NOT NULL,
                usuario TEXT UNIQUE NOT NULL,
                password TEXT NOT NULL,
                rol TEXT NOT NULL DEFAULT 'operador',
                activo INTEGER NOT NULL DEFAULT 1,
                fecha_creacion TEXT NOT NULL
            )
        ''')

        # Usuario administrador inicial (Maicol) si no existe
        cur = conn.execute("SELECT COUNT(*) FROM usuarios WHERE usuario = 'maicol'")
        if cur.fetchone()[0] == 0:
            ahora = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            conn.execute('''
                INSERT INTO usuarios (nombre, usuario, password, rol, activo, fecha_creacion)
                VALUES ('Maicol', 'maicol', 'admin123', 'admin', 1, ?)
            ''', (ahora,))

        # Tabla de registros de motos
        conn.execute('''
            CREATE TABLE IF NOT EXISTS registros (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                fecha_registro TEXT NOT NULL,
                mes_periodo TEXT NOT NULL,
                placa TEXT NOT NULL,
                cedula TEXT,
                nombre_cliente TEXT NOT NULL,
                telefono TEXT,
                valor_pagado REAL NOT NULL DEFAULT 0,
                concesionario TEXT NOT NULL,
                comision REAL NOT NULL DEFAULT 0,
                banco TEXT,
                numero_referencia TEXT,
                archivo_comprobante TEXT,
                observaciones TEXT,
                creado_por TEXT DEFAULT 'maicol'
            )
        ''')
        
        # Migración automática si falta creado_por
        cur = conn.execute("PRAGMA table_info(registros)")
        cols = [r["name"] for r in cur.fetchall()]
        if cols and "creado_por" not in cols:
            conn.execute("ALTER TABLE registros ADD COLUMN creado_por TEXT DEFAULT 'maicol'")

        conn.commit()

init_db()

def obtener_usuario_actual(req):
    """Verifica credenciales enviadas en encabezados HTTP o query parameters."""
    auth_user = req.headers.get("X-User-Name") or req.args.get("u")
    auth_pass = req.headers.get("X-User-Pass") or req.args.get("p")
    
    # Si viene desde localhost y no envió usuario, por comodidad en su propia máquina tiene acceso Maicol
    if not auth_user and (req.remote_addr in ['127.0.0.1', '::1', 'localhost']):
        with get_db() as conn:
            cur = conn.execute("SELECT * FROM usuarios WHERE usuario = 'maicol' AND activo = 1")
            row = cur.fetchone()
            if row:
                return dict(row)

    if not auth_user or not auth_pass:
        return None

    with get_db() as conn:
        cur = conn.execute("SELECT * FROM usuarios WHERE usuario = ? AND password = ? AND activo = 1", (auth_user, auth_pass))
        row = cur.fetchone()
        return dict(row) if row else None

def es_admin(req):
    u = obtener_usuario_actual(req)
    return u is not None and u.get("rol") == "admin"

def puede_registrar(req):
    u = obtener_usuario_actual(req)
    return u is not None and u.get("rol") in ["admin", "operador"]

def extraer_datos_pdf(pdf_path):
    """Extrae texto del PDF de PSE e intenta predecir banco, valor pagado y referencia."""
    texto = ""
    try:
        reader = PdfReader(pdf_path)
        for page in reader.pages:
            t = page.extract_text()
            if t:
                texto += t + "\n"
    except Exception as e:
        print(f"Error leyendo PDF: {e}")

    sugerencias = {
        "banco": "",
        "valor_pagado": None,
        "referencia": "",
        "texto_detectado": texto[:800]
    }

    texto_lower = texto.lower()

    if "bancolombia" in texto_lower:
        sugerencias["banco"] = "Bancolombia (PSE)"
    elif "bogota" in texto_lower or "bogotá" in texto_lower:
        sugerencias["banco"] = "Banco de Bogotá (PSE)"
    elif "davivienda" in texto_lower:
        sugerencias["banco"] = "Davivienda (PSE)"
    elif "pse" in texto_lower:
        sugerencias["banco"] = "PSE en línea"

    ref_match = re.search(r'(?:cus|referencia|autorizaci[oó]n|transacci[oó]n|recibo|aprobaci[oó]n)[:\s#]*([A-Za-z0-9\-_]{5,20})', texto, re.IGNORECASE)
    if ref_match:
        sugerencias["referencia"] = ref_match.group(1).strip()

    patrones_valor = [
        r'\$\s*([0-9]{1,3}(?:[.,][0-9]{3})+(?:[.,][0-9]{2})?)',
        r'valor[:\s]*\$?\s*([0-9]{1,3}(?:[.,][0-9]{3})+(?:[.,][0-9]{2})?)',
        r'total[:\s]*\$?\s*([0-9]{1,3}(?:[.,][0-9]{3})+(?:[.,][0-9]{2})?)',
    ]

    for pat in patrones_valor:
        matches = re.finditer(pat, texto, re.IGNORECASE)
        for m in matches:
            val_str = m.group(1).replace(' ', '')
            if ',' in val_str and '.' in val_str:
                if val_str.rfind(',') > val_str.rfind('.'):
                    val_str = val_str.replace('.', '').replace(',', '.')
                else:
                    val_str = val_str.replace(',', '')
            elif '.' in val_str and len(val_str.split('.')[-1]) == 3:
                val_str = val_str.replace('.', '')
            elif ',' in val_str and len(val_str.split(',')[-1]) == 3:
                val_str = val_str.replace(',', '')
            
            try:
                val_float = float(val_str)
                if 100000 <= val_float <= 15000000:
                    sugerencias["valor_pagado"] = val_float
                    break
            except ValueError:
                continue
        if sugerencias["valor_pagado"]:
            break

    placa_match = re.search(r'\b([A-Z]{3}[0-9]{2}[A-Z]|[A-Z]{3}[0-9]{3})\b', texto.upper())
    if placa_match:
        sugerencias["placa"] = placa_match.group(1)

    return sugerencias

@app.route("/")
def index():
    return render_template("index.html")

# ================= RUTAS DE AUTENTICACIÓN Y USUARIOS =================

@app.route("/api/login", methods=["POST"])
def login():
    data = request.get_json() or {}
    user = str(data.get("usuario", "")).strip().lower()
    pwd = str(data.get("password", "")).strip()

    with get_db() as conn:
        cur = conn.execute("SELECT id, nombre, usuario, rol FROM usuarios WHERE usuario = ? AND password = ? AND activo = 1", (user, pwd))
        row = cur.fetchone()
        if row:
            return jsonify({"success": True, "usuario": dict(row)})
    return jsonify({"success": False, "error": "Usuario o contraseña incorrectos"}), 401

@app.route("/api/usuarios", methods=["GET"])
def listar_usuarios():
    if not es_admin(request):
        return jsonify({"error": "Solo el administrador puede ver los usuarios"}), 403

    with get_db() as conn:
        cur = conn.execute("SELECT id, nombre, usuario, rol, activo, fecha_creacion FROM usuarios ORDER BY id ASC")
        filas = [dict(r) for r in cur.fetchall()]
    return jsonify({"usuarios": filas})

@app.route("/api/usuarios", methods=["POST"])
def crear_usuario():
    if not es_admin(request):
        return jsonify({"error": "Solo el administrador puede crear usuarios"}), 403

    data = request.get_json() or {}
    nombre = str(data.get("nombre", "")).strip()
    user = str(data.get("usuario", "")).strip().lower()
    pwd = str(data.get("password", "")).strip()
    rol = str(data.get("rol", "operador")).strip().lower()

    if not nombre or not user or not pwd:
        return jsonify({"error": "Nombre, usuario y contraseña son obligatorios"}), 400

    if rol not in ["admin", "operador", "consulta"]:
        rol = "operador"

    ahora = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    try:
        with get_db() as conn:
            conn.execute('''
                INSERT INTO usuarios (nombre, usuario, password, rol, activo, fecha_creacion)
                VALUES (?, ?, ?, ?, 1, ?)
            ''', (nombre, user, pwd, rol, ahora))
            conn.commit()
    except sqlite3.IntegrityError:
        return jsonify({"error": f"El nombre de usuario '{user}' ya existe"}), 400

    return jsonify({"success": True})

@app.route("/api/usuarios/<int:id>", methods=["DELETE"])
def eliminar_usuario(id):
    if not es_admin(request):
        return jsonify({"error": "Solo el administrador puede eliminar usuarios"}), 403

    usuario_actual = obtener_usuario_actual(request)
    if usuario_actual and usuario_actual.get("id") == id:
        return jsonify({"error": "No puedes eliminar tu propia cuenta de administrador"}), 400

    with get_db() as conn:
        conn.execute("DELETE FROM usuarios WHERE id = ?", (id,))
        conn.commit()

    return jsonify({"success": True})

# ================= RUTAS DE MANTENIMIENTO =================

@app.route("/api/vaciar-historial", methods=["POST"])
def vaciar_historial():
    """Elimina todos los registros y comprobantes físicos (Solo Admin con confirmación de clave)."""
    if not es_admin(request):
        return jsonify({"error": "Solo el administrador puede vaciar el historial"}), 403

    data = request.get_json() or {}
    pwd_confirm = str(data.get("password", "")).strip()

    usuario_actual = obtener_usuario_actual(request)
    if not usuario_actual or usuario_actual.get("password") != pwd_confirm:
        return jsonify({"error": "Contraseña de confirmación incorrecta"}), 401

    # 1. Vaciar tabla SQLite
    with get_db() as conn:
        conn.execute("DELETE FROM registros")
        conn.execute("DELETE FROM sqlite_sequence WHERE name='registros'")
        conn.commit()

    # 2. Vaciar carpeta de comprobantes
    try:
        if COMPROBANTES_DIR.exists():
            for item in COMPROBANTES_DIR.iterdir():
                if item.is_dir():
                    shutil.rmtree(item, ignore_errors=True)
                elif item.is_file():
                    item.unlink(missing_ok=True)
    except Exception as e:
        print(f"Error limpiando comprobantes: {e}")

    return jsonify({"success": True, "mensaje": "Historial y comprobantes vaciados por completo"})

# ================= RUTAS DE REGISTROS Y ARCHIVOS =================

@app.route("/api/info-conexion", methods=["GET"])
def info_conexion():
    return jsonify({
        "enlace_publico": ENLACE_PUBLICO_ACTUAL,
        "ip_local": obtener_ip_local(),
        "hostname": socket.gethostname()
    })

@app.route("/api/analizar-pdf", methods=["POST"])
def analizar_pdf():
    if "archivo" not in request.files:
        return jsonify({"error": "No se envió ningún archivo"}), 400
    
    archivo = request.files["archivo"]
    if archivo.filename == "":
        return jsonify({"error": "Archivo sin nombre"}), 400

    temp_dir = BASE_DIR / "temp"
    temp_dir.mkdir(exist_ok=True)
    temp_path = temp_dir / secure_filename(f"temp_{datetime.datetime.now().timestamp()}_{archivo.filename}")
    archivo.save(temp_path)

    try:
        datos = extraer_datos_pdf(temp_path)
    finally:
        if temp_path.exists():
            try:
                temp_path.unlink()
            except Exception:
                pass

    return jsonify({"sugerencias": datos})

@app.route("/api/registros", methods=["GET"])
def listar_registros():
    q = request.args.get("q", "").strip()
    mes = request.args.get("mes", "").strip()
    concesionario = request.args.get("concesionario", "").strip()
    admin_auth = es_admin(request)

    sql = "SELECT * FROM registros WHERE 1=1"
    params = []

    if q:
        sql += " AND (placa LIKE ? OR cedula LIKE ? OR nombre_cliente LIKE ? OR telefono LIKE ? OR numero_referencia LIKE ?)"
        wild = f"%{q}%"
        params.extend([wild, wild, wild, wild, wild])

    if mes:
        sql += " AND mes_periodo = ?"
        params.append(mes)

    if concesionario:
        sql += " AND concesionario = ?"
        params.append(concesionario)

    sql += " ORDER BY id DESC"

    with get_db() as conn:
        cursor = conn.execute(sql, params)
        rows = [dict(row) for row in cursor.fetchall()]

    if not admin_auth:
        for r in rows:
            r["comision"] = None  # Protección de privacidad contable

    return jsonify({"registros": rows, "es_admin": admin_auth})

@app.route("/api/registros", methods=["POST"])
def crear_registro():
    if not puede_registrar(request):
        return jsonify({"error": "No tienes permisos para registrar motos. Inicia sesión con una cuenta autorizada."}), 403

    usuario_actual = obtener_usuario_actual(request)
    creado_por = usuario_actual.get("usuario") if usuario_actual else "maicol"

    placa = request.form.get("placa", "").strip().upper()
    cedula = request.form.get("cedula", "").strip()
    nombre_cliente = request.form.get("nombre_cliente", "").strip().title()
    telefono = request.form.get("telefono", "").strip()
    concesionario = request.form.get("concesionario", "").strip()
    banco = request.form.get("banco", "").strip()
    numero_referencia = request.form.get("numero_referencia", "").strip()
    observaciones = request.form.get("observaciones", "").strip()
    
    try:
        valor_pagado = float(request.form.get("valor_pagado", 0))
    except (ValueError, TypeError):
        valor_pagado = 0.0

    try:
        comision = float(request.form.get("comision", 0))
    except (ValueError, TypeError):
        comision = 0.0

    if not placa or not nombre_cliente or not concesionario:
        return jsonify({"error": "Placa, Nombre y Concesionario son obligatorios"}), 400

    ahora = datetime.datetime.now()
    mes_periodo = ahora.strftime("%Y-%m")
    fecha_registro = ahora.strftime("%Y-%m-%d %H:%M:%S")

    archivo_relativo = None
    if "archivo" in request.files:
        archivo = request.files["archivo"]
        if archivo and archivo.filename != "":
            carpeta_mes = COMPROBANTES_DIR / mes_periodo
            carpeta_mes.mkdir(parents=True, exist_ok=True)

            placa_limpia = re.sub(r'[^A-Za-z0-9]', '', placa)
            cedula_limpia = re.sub(r'[^0-9]', '', cedula) or "SC"
            timestamp_corto = ahora.strftime("%H%M%S")
            nombre_guardado = f"{mes_periodo}_{placa_limpia}_{cedula_limpia}_{timestamp_corto}.pdf"
            
            ruta_destino = carpeta_mes / nombre_guardado
            archivo.save(ruta_destino)
            archivo_relativo = f"{mes_periodo}/{nombre_guardado}"

    with get_db() as conn:
        cursor = conn.execute('''
            INSERT INTO registros (
                fecha_registro, mes_periodo, placa, cedula, nombre_cliente,
                telefono, valor_pagado, concesionario, comision, banco,
                numero_referencia, archivo_comprobante, observaciones, creado_por
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (
            fecha_registro, mes_periodo, placa, cedula, nombre_cliente,
            telefono, valor_pagado, concesionario, comision, banco,
            numero_referencia, archivo_relativo, observaciones, creado_por
        ))
        conn.commit()
        nuevo_id = cursor.lastrowid

    return jsonify({"success": True, "id": nuevo_id, "archivo": archivo_relativo})

@app.route("/api/registros/<int:id>", methods=["DELETE"])
def eliminar_registro(id):
    if not es_admin(request):
        return jsonify({"error": "Solo el administrador puede eliminar registros individuales"}), 403

    with get_db() as conn:
        cursor = conn.execute("SELECT archivo_comprobante FROM registros WHERE id = ?", (id,))
        row = cursor.fetchone()
        if not row:
            return jsonify({"error": "Registro no encontrado"}), 404

        if row["archivo_comprobante"]:
            archivo_path = COMPROBANTES_DIR / row["archivo_comprobante"]
            if archivo_path.exists():
                try:
                    archivo_path.unlink()
                except Exception as e:
                    print(f"Error eliminando archivo: {e}")

        conn.execute("DELETE FROM registros WHERE id = ?", (id,))
        conn.commit()

    return jsonify({"success": True})

@app.route("/api/metricas", methods=["GET"])
def obtener_metricas():
    mes = request.args.get("mes", "").strip() or datetime.datetime.now().strftime("%Y-%m")
    admin_auth = es_admin(request)

    with get_db() as conn:
        cur = conn.execute('''
            SELECT 
                COUNT(*) as total_motos,
                COALESCE(SUM(valor_pagado), 0) as total_pagado_seguros,
                COALESCE(SUM(comision), 0) as total_comisiones,
                COALESCE(SUM(valor_pagado + comision), 0) as total_ingresado_cuenta
            FROM registros
            WHERE mes_periodo = ?
        ''', (mes,))
        metricas_mes = dict(cur.fetchone())

        cur_global = conn.execute('''
            SELECT 
                COUNT(*) as total_motos,
                COALESCE(SUM(valor_pagado), 0) as total_pagado_seguros,
                COALESCE(SUM(comision), 0) as total_comisiones,
                COALESCE(SUM(valor_pagado + comision), 0) as total_ingresado_cuenta
            FROM registros
        ''')
        metricas_global = dict(cur_global.fetchone())

        cur_meses = conn.execute('SELECT DISTINCT mes_periodo FROM registros ORDER BY mes_periodo DESC')
        meses = [r[0] for r in cur_meses.fetchall()]
        if mes not in meses:
            meses.insert(0, mes)

        cur_conc = conn.execute('SELECT DISTINCT concesionario FROM registros WHERE concesionario IS NOT NULL AND concesionario != "" ORDER BY concesionario ASC')
        concesionarios = [r[0] for r in cur_conc.fetchall()]

    if not admin_auth:
        metricas_mes["total_comisiones"] = None
        metricas_mes["total_ingresado_cuenta"] = None
        metricas_global["total_comisiones"] = None
        metricas_global["total_ingresado_cuenta"] = None

    return jsonify({
        "es_admin": admin_auth,
        "mes_seleccionado": mes,
        "mes": metricas_mes,
        "global": metricas_global,
        "meses_disponibles": meses,
        "concesionarios": concesionarios
    })

@app.route("/comprobantes/<path:filepath>")
def servir_comprobante(filepath):
    return send_from_directory(COMPROBANTES_DIR, filepath)

@app.route("/api/exportar-excel", methods=["GET"])
def exportar_excel():
    if not es_admin(request):
        return jsonify({"error": "Solo el administrador puede descargar el reporte fiscal"}), 403

    mes = request.args.get("mes", "").strip()

    sql = "SELECT * FROM registros"
    params = []
    if mes:
        sql += " WHERE mes_periodo = ?"
        params.append(mes)
    sql += " ORDER BY fecha_registro ASC"

    with get_db() as conn:
        cursor = conn.execute(sql, params)
        filas = cursor.fetchall()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Soportes DIAN - {mes or 'Historico'}"
    ws.views.sheetView[0].showGridLines = True

    fuente_titulo = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    fuente_subtitulo = Font(name="Calibri", size=11, italic=True, color="4B5563")
    fuente_encabezado = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    relleno_encabezado = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    
    fuente_totales = Font(name="Calibri", size=11, bold=True, color="111827")
    relleno_totales = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")

    borde_fino = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB")
    )

    ws["A1"] = "RELACIÓN DE SOPORTES Y PAGO DE SEGUROS MG (INGRESOS PARA TERCEROS)"
    ws["A1"].font = fuente_titulo
    ws["A2"] = f"Periodo: {mes or 'HISTÓRICO COMPLETO'} | Generado: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')} | Justificación Contable ante la DIAN"
    ws["A2"].font = fuente_subtitulo

    encabezados = [
        "ID", "Fecha / Hora", "Mes Periodo", "Placa", "Cédula Cliente",
        "Nombre del Comprador", "Teléfono", "Concesionario", "Banco / PSE",
        "No. Referencia CUS", "Valor Pagado Seguro (Terceros)",
        "Comisión Ganada (Ingreso Propio)", "Total Entrado a Cuenta", "Comprobante PDF"
    ]

    for col_idx, texto in enumerate(encabezados, start=1):
        celda = ws.cell(row=4, column=col_idx, value=texto)
        celda.font = fuente_encabezado
        celda.fill = relleno_encabezado
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[4].height = 28
    fila_inicio = 5
    fila_actual = fila_inicio

    for f in filas:
        ws.cell(row=fila_actual, column=1, value=f["id"]).alignment = Alignment(horizontal="center")
        ws.cell(row=fila_actual, column=2, value=f["fecha_registro"]).alignment = Alignment(horizontal="center")
        ws.cell(row=fila_actual, column=3, value=f["mes_periodo"]).alignment = Alignment(horizontal="center")
        ws.cell(row=fila_actual, column=4, value=f["placa"]).alignment = Alignment(horizontal="center")
        ws.cell(row=fila_actual, column=5, value=f["cedula"] or "N/A").alignment = Alignment(horizontal="center")
        ws.cell(row=fila_actual, column=6, value=f["nombre_cliente"])
        ws.cell(row=fila_actual, column=7, value=f["telefono"] or "N/A").alignment = Alignment(horizontal="center")
        ws.cell(row=fila_actual, column=8, value=f["concesionario"])
        ws.cell(row=fila_actual, column=9, value=f["banco"] or "PSE")
        ws.cell(row=fila_actual, column=10, value=f["numero_referencia"] or "S/R").alignment = Alignment(horizontal="center")
        
        c_val = ws.cell(row=fila_actual, column=11, value=f["valor_pagado"])
        c_val.number_format = '$ #,##0'
        
        c_com = ws.cell(row=fila_actual, column=12, value=f["comision"])
        c_com.number_format = '$ #,##0'
        
        c_tot = ws.cell(row=fila_actual, column=13, value=f"=K{fila_actual}+L{fila_actual}")
        c_tot.number_format = '$ #,##0'

        celda_link = ws.cell(row=fila_actual, column=14)
        if f["archivo_comprobante"]:
            ruta_pdf = COMPROBANTES_DIR / f["archivo_comprobante"]
            if ruta_pdf.exists():
                uri_pdf = ruta_pdf.resolve().as_uri()
                # Fórmula HYPERLINK nativa de Excel para que Windows lo abra directamente sin error
                celda_link.value = f'=HYPERLINK("{uri_pdf}", "Abrir PDF ({f["placa"]})")'
                celda_link.font = Font(name="Calibri", color="2563EB", underline="single")
                celda_link.alignment = Alignment(horizontal="center")
            else:
                celda_link.value = "Archivo no encontrado"
                celda_link.alignment = Alignment(horizontal="center")
        else:
            celda_link.value = "Sin Comprobante"
            celda_link.alignment = Alignment(horizontal="center")

        for col_idx in range(1, 15):
            ws.cell(row=fila_actual, column=col_idx).border = borde_fino

        fila_actual += 1

    if fila_actual > fila_inicio:
        fila_totales = fila_actual
        ws.cell(row=fila_totales, column=1, value="TOTALES:").font = fuente_totales
        ws.merge_cells(start_row=fila_totales, start_column=1, end_row=fila_totales, end_column=10)
        ws.cell(row=fila_totales, column=1).alignment = Alignment(horizontal="right")
        
        c_tot_val = ws.cell(row=fila_totales, column=11, value=f"=SUM(K{fila_inicio}:K{fila_totales-1})")
        c_tot_val.font = fuente_totales
        c_tot_val.number_format = '$ #,##0'

        c_tot_com = ws.cell(row=fila_totales, column=12, value=f"=SUM(L{fila_inicio}:L{fila_totales-1})")
        c_tot_com.font = fuente_totales
        c_tot_com.number_format = '$ #,##0'

        c_tot_tot = ws.cell(row=fila_totales, column=13, value=f"=SUM(M{fila_inicio}:M{fila_totales-1})")
        c_tot_tot.font = fuente_totales
        c_tot_tot.number_format = '$ #,##0'

        for col_idx in range(1, 15):
            celda = ws.cell(row=fila_totales, column=col_idx)
            celda.fill = relleno_totales
            celda.border = borde_fino

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in [1, 2]:
                continue
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    temp_export = BASE_DIR / f"Reporte_DIAN_Seguros_MG_{mes or 'Historico'}.xlsx"
    wb.save(temp_export)

    return send_file(
        temp_export,
        as_attachment=True,
        download_name=f"Reporte_DIAN_Seguros_MG_{mes or 'Historico'}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

def obtener_ip_local():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except Exception:
        return "127.0.0.1"

def iniciar_localtunnel_fijo():
    """Inicia localtunnel con subdominio personalizado transito-mg."""
    global ENLACE_PUBLICO_ACTUAL
    try:
        proc = subprocess.Popen(
            ["npx", "--yes", "localtunnel", "--port", "5000", "--subdomain", "transito-mg"],
            shell=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            bufsize=1
        )
        for line in proc.stdout:
            if "your url is:" in line.lower() or "loca.lt" in line.lower():
                m = re.search(r'https://[a-zA-Z0-9-]+\.loca\.lt', line)
                if m:
                    ENLACE_PUBLICO_ACTUAL = m.group(0)
                    print("\n" + "=" * 70)
                    print("🌐 ¡DOMINIO PROPIO FIJO ACTIVADO (SIN PAGAR)!")
                    print(f"👉 TU ENLACE FIJO: {ENLACE_PUBLICO_ACTUAL}")
                    print("=" * 70 + "\n")
                    with open(BASE_DIR / "ENLACE_PARA_COMPARTIR.txt", "w", encoding="utf-8") as f:
                        f.write(f"DOMINIO SEGURO FIJO:\n{ENLACE_PUBLICO_ACTUAL}\n\n(Abre este enlace desde cualquier computador o celular)\n")
                    break
    except Exception as e:
        print(f"Error iniciando túnel con dominio propio: {e}")

if __name__ == "__main__":
    ip_red = obtener_ip_local()
    hostname = socket.gethostname()
    compartir_web = "--compartir" in sys.argv

    if compartir_web:
        hilo = threading.Thread(target=iniciar_localtunnel_fijo, daemon=True)
        hilo.start()

    print("=" * 65)
    print("🚀 SOPORTES Y SEGUROS TRÁNSITO MG - SISTEMA INICIADO")
    print(f"👉 En tu computador (Local):      http://127.0.0.1:5000")
    print(f"👉 En red Wi-Fi (Sin números):    http://{hostname}:5000")
    print(f"👉 En red Wi-Fi (IP directa):     http://{ip_red}:5000")
    if compartir_web:
        print("⏳ Conectando con tu dominio propio gratis: https://transito-mg.loca.lt ...")
    print("=" * 65)
    
    app.run(host="0.0.0.0", port=5000, debug=False)
